from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

# Slide 1: Title
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "MusicTalk Platform"
subtitle.text = "A Microservice Approach to AI-Driven Musical Instrument Detection\n\nArchitecture & Implementation Overview"

# Slide 2: Introduction
bullet_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
title_shape.text = "What is MusicTalk?"
body_shape = shapes.placeholders[1]
tf = body_shape.text_frame
tf.text = "An intelligent ecosystem for audio analysis and music discovery."
p = tf.add_paragraph()
p.text = "Microservice-Based Architecture: Highly available, modular ecosystem."
p.level = 1
p = tf.add_paragraph()
p.text = "Deep Learning Audio Recognition: Analyzes raw audio via CNN models to detect instruments."
p.level = 1
p = tf.add_paragraph()
p.text = "Smart Catalog Engine: Persists acoustic structures and overlaps natively in MySQL."
p.level = 1
p = tf.add_paragraph()
p.text = "Continuous Auto-Play Context: Dynamically recommends and plays similar tracks via Jaccard Overlap ratios."
p.level = 1

# Slide 3: Microservice Architecture
slide = prs.slides.add_slide(bullet_slide_layout)
title_shape = slide.shapes.title
title_shape.text = "Microservice Ecosystem"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "Seamless integration across decoupled services:"
p = tf.add_paragraph()
p.text = "Frontend App: React (Vite) containing Dashboards, Audio Analyzer, and Auto-DJ Player."
p.level = 1
p = tf.add_paragraph()
p.text = "API Gateway (FastAPI): Reverse proxies, secures, and routes UI requests."
p.level = 1
p = tf.add_paragraph()
p.text = "Recognition Service (PyTorch): In-memory CNN evaluator & Librosa spectrogram generator."
p.level = 1
p = tf.add_paragraph()
p.text = "Catalog Service (SQLAlchemy): Core CRUD processor and disk-storage coordinator."
p.level = 1
p = tf.add_paragraph()
p.text = "Recommendation Engine: Computes real-time discovery based on instrument statistical maps."
p.level = 1

# Slide 4: Database Schema (songs table)
slide = prs.slides.add_slide(bullet_slide_layout)
title_shape = slide.shapes.title
title_shape.text = "Database Design - 'songs' Table"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "Adopts a hybrid approach: Heavy metadata in MySQL, binary audio on Disk."
p = tf.add_paragraph()
p.text = "ID & Metadata: id, title, artist, album, genre, duration_sec"
p.level = 1
p = tf.add_paragraph()
p.text = "ML Embeddings: instruments (JSON), tags (JSON)"
p.level = 1
p = tf.add_paragraph()
p.text = "Highlights: We use schema-less JSON to capture dynamic ML model outputs like {confidence, overlap}."
p.level = 2
p = tf.add_paragraph()
p.text = "Resource Locators: file_hash, storage_path, audio_url, cover_url"
p.level = 1
p = tf.add_paragraph()
p.text = "System: status (UPLOADED, DONE, etc.), created_at, updated_at"
p.level = 1

# Slide 5: Machine Learning & Recognition Flow
slide = prs.slides.add_slide(bullet_slide_layout)
title_shape = slide.shapes.title
title_shape.text = "Machine Learning Pipeline"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "How an audio file is translated into acoustic embeddings:"
p = tf.add_paragraph()
p.text = "Step 1: ID3 Parsing (Mutagen) & Audio Loading (Librosa)."
p.level = 1
p = tf.add_paragraph()
p.text = "Step 2: Generation of short-time Log-Mel Spectrograms (3.0s window, 1.5s hop)."
p.level = 1
p = tf.add_paragraph()
p.text = "Step 3: InstrumentAggregator streams spectrograms through iterative PyTorch CNNs."
p.level = 1
p = tf.add_paragraph()
p.text = "Step 4: Sigmoid activation constructs time-series confidence graphs mapped to JSON outputs."
p.level = 1

# Slide 6: Deep Tech Workflow Summary
slide = prs.slides.add_slide(bullet_slide_layout)
title_shape = slide.shapes.title
title_shape.text = "Workflow & Scalability Summary"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "Delivering a scalable streaming and AI platform:"
p = tf.add_paragraph()
p.text = "Efficient Persistence: BLOB processing kept out of MySQL; relying strictly on metadata mappings."
p.level = 1
p = tf.add_paragraph()
p.text = "Dashboard Real-Time Aggregation: Analyzes JSON schemas across catalogs dynamically without static caches."
p.level = 1
p = tf.add_paragraph()
p.text = "Robust Streaming Lifecycle: File > API Gateway > PyTorch Map > MySQL Record > Jaccard Recommendation Array > Auto Play UI."
p.level = 1

prs.save("C:/204/MusicTalk_Presentation.pptx")
print("Presentation successfully saved to C:/204/MusicTalk_Presentation.pptx")
